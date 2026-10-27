# -*- coding: utf-8 -*-
"""
Excel(.xlsx) 题库融合解析器

将以下两个来源解析、去重后融合为一个题库：

    data/光伏发电运维值班员题库.xlsx   （结构完整）
    data/光伏题库全.xlsx               （结构化 + 半结构化 + OCR 文本）

输出统一字段（与 static_version/data/*.json 一致）：
    id / type / difficulty / question
    单选/多选: options, answer
    判断:      options(√,×), answer
    填空:      answers, orderFixed
    简答/计算/论述/绘图: answer, images / answerImages / answerMissing

说明：
- OCR/半结构化部分采用尽力解析（best-effort），解析失败或信息不足的条目会跳过并记入 warnings。
- 同一题库内按“题干 + 选项 + 答案”去重；题干相同但答案/选项不同的变体保留。
- 图片按哈希命名，返回 {文件名: bytes}，由 build 脚本落盘。
"""
from __future__ import annotations

import hashlib
import os
import re
import zipfile
from collections import OrderedDict
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import load_workbook

from qbank.docx_parser import mathify

# ---------------------------------------------------------------------------
# 通用工具
# ---------------------------------------------------------------------------

LEVEL_MAP = {
    "B1": "一级生产岗位能力",
    "B2": "二级生产岗位能力",
    "B3": "三级生产岗位能力",
    "B4": "四级生产岗位能力",
    "B5": "五级生产岗位能力",
}

_CALC = re.compile(
    r"(计算|试求|求出|试算|求其|求该|求这|求一|多少|几块|几台|几路|几串|几根|"
    r"偏差|利用小时|运行指标|上网电量|(?<![要需请诉])求)"
)

_SPACES = re.compile(r"[ \t\u3000]+")
_NUM_PREFIX = re.compile(r"^[\s\u3000]*[0-9０-９]+\s*[.．、,，:：]\s*")
_PAGE = re.compile(r"^\s*\d+\s*/\s*\d+\s*[PＰ]?\s*$")
_DIFF = {"易": "容易", "较易": "容易", "中": "中等", "中等": "中等",
         "难": "困难", "较难": "困难", "容易": "容易", "困难": "困难"}
_HEADERS = {"单选题", "多选题", "判断题", "填空题", "简答题", "论述题", "计算题",
            "绘图题", "问答题", "设备实操"}


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else ("%g" % v)
    if isinstance(v, int):
        return str(v)
    return str(v).strip()


def _text(v) -> str:
    raw = _s(v).replace("\r\n", "\n").replace("\r", "\n")
    lines = [ln.strip() for ln in raw.split("\n")]
    return "\n".join(ln for ln in lines if ln)


def _norm(s: str) -> str:
    """题干/答案归一化，用于去重。"""
    s = _s(s)
    s = re.sub(r"\s+", "", s)
    s = re.sub(r"[，。、；：？！,.;:?!\"'“”‘’（）()《》〈〉【】\[\]]+", "", s)
    return s.lower()


def _diff(raw: str) -> str:
    return _DIFF.get(_s(raw), "")


def _clean_question(q: str) -> str:
    q = _NUM_PREFIX.sub("", _text(q)).strip()
    q = re.sub(r"[ \t\u3000]+", " ", q)
    return q.strip()


def _maybe_mathify(s: str, qtype: str) -> str:
    """计算/绘图题保留原始文字，不做 LaTeX 转换；其余题型沿用 mathify。"""
    if qtype in ("计算", "绘图"):
        return _s(s)
    return mathify(s)


def _fix_subjective_qa(q: str, a: str) -> Tuple[str, str]:
    """修正简答/论述题中“答：”与题干粘连、答案自带“答：”前缀的问题。"""
    q = _s(q)
    a = _s(a)
    m = re.search(r"(?<!回)答\s*[：:；;]\s*", q)
    if m:
        rest = q[m.end():].strip()
        q = q[:m.start()].strip()
        if rest:
            a = (rest + "\n" + a).strip()
    a = re.sub(r"^\s*答\s*[：:；;]\s*", "", a).strip()
    return q, a


def _answer_letter(s: str) -> str:
    letters = "".join(sorted(set(_s(s).upper()) & set("ABCDEFGHIJKL")))
    return letters


def _norm_judge(v: str) -> str:
    s = _s(v).replace("|", "").strip()
    s = s.replace("\ufe0f", "").replace("\u200d", "")
    if s in ("√", "✔", "✓", "✅", "对", "正确", "是", "T", "t", "true", "TRUE"):
        return "√"
    if s in ("×", "x", "X", "✗", "✘", "❌", "❎", "错", "错误", "否", "F", "f", "false", "FALSE"):
        return "×"
    return ""


# ---------------------------------------------------------------------------
# 图片
# ---------------------------------------------------------------------------

def _collect_images(ws) -> Tuple[Dict[str, bytes], Dict[int, List[str]]]:
    blobs: Dict[str, bytes] = {}
    anchors: Dict[int, List[str]] = {}
    for im in getattr(ws, "_images", []) or []:
        try:
            data = im._data()
        except Exception:
            try:
                data = im.ref.read()
            except Exception:
                continue
        ext = (getattr(im, "format", None) or "png").lower()
        if ext == "jpg":
            ext = "jpeg"
        h = hashlib.md5(data).hexdigest()[:16]
        fname = f"img_{h}.{ext}"
        blobs[fname] = data
        try:
            row = im.anchor._from.row + 1
        except Exception:
            continue
        anchors.setdefault(row, []).append(fname)
    return blobs, anchors


def _collect_cell_images(path: str) -> Tuple[Dict[str, str], Dict[str, bytes]]:
    """
    读取 WPS 单元格图片（cellimages.xml），返回：
        {DISPIMG_ID: 文件名}
        {文件名: bytes}
    用于处理 `=DISPIMG("ID_...",1)` 形式的答案。
    """
    cell_map: Dict[str, str] = {}
    blobs: Dict[str, bytes] = {}
    try:
        zf = zipfile.ZipFile(path)
    except Exception:
        return cell_map, blobs
    with zf:
        names = set(zf.namelist())
        if "xl/cellimages.xml" not in names or "xl/_rels/cellimages.xml.rels" not in names:
            return cell_map, blobs
        try:
            xml = zf.read("xl/cellimages.xml").decode("utf-8", "ignore")
            rels = zf.read("xl/_rels/cellimages.xml.rels").decode("utf-8", "ignore")
        except Exception:
            return cell_map, blobs
        rel_map: Dict[str, str] = {}
        for m in re.finditer(r"<Relationship\b[^>]*>", rels):
            tag = m.group(0)
            rid = re.search(r'Id="([^"]+)"', tag)
            target = re.search(r'Target="([^"]+)"', tag)
            if rid and target:
                rel_map[rid.group(1)] = target.group(1)
        for m in re.finditer(
            r'<xdr:cNvPr[^>]*name="([^"]+)"[^>]*>.*?<a:blip[^>]*r:embed="([^"]+)"',
            xml, re.S,
        ):
            img_id, rid = m.group(1), m.group(2)
            target = rel_map.get(rid)
            if not target:
                continue
            zpath = "xl/" + target.lstrip("/")
            if zpath not in names:
                continue
            data = zf.read(zpath)
            ext = os.path.splitext(target)[1].lower().lstrip(".") or "png"
            if ext == "jpg":
                ext = "jpeg"
            h = hashlib.md5(data).hexdigest()[:16]
            fname = f"img_{h}.{ext}"
            cell_map[img_id] = fname
            blobs[fname] = data
    return cell_map, blobs


# ---------------------------------------------------------------------------
# 值班员题库（结构完整）
# ---------------------------------------------------------------------------

_DUTY_SHEETS = ("单选题", "多选题", "判断题", "填空题", "问答题")


def _parse_duty_objective(ws, qtype: str) -> List[Dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [_s(x) for x in rows[0]]
    try:
        ai = next(i for i, x in enumerate(header) if x.startswith("正确答案"))
    except StopIteration:
        return []
    diff_i = header.index("难度：") if "难度：" in header else 2
    opt_i = header.index("选项A") if "选项A" in header else 5
    out: List[Dict] = []
    for row in rows[1:]:
        q = _text(row[0]) if len(row) > 0 else ""
        if not q:
            continue
        raw_ans = _s(row[ai]) if len(row) > ai else ""
        diff = _diff(row[diff_i]) if len(row) > diff_i else ""
        if qtype == "判断":
            ans = _norm_judge(raw_ans)
            if not ans:
                continue
            out.append({"type": "判断", "difficulty": diff,
                        "question": mathify(_clean_question(q)),
                        "options": ["√", "×"], "answer": ans})
            continue

        ans = _answer_letter(raw_ans)
        if not ans:
            continue
        opts = [mathify(_text(row[i])) for i in range(opt_i, len(row))
                if i < len(row) and _s(row[i])]
        if len(opts) < 2:
            continue
        # 去掉答案中超出选项范围的字母
        ans = "".join(c for c in ans if (ord(c) - 65) < len(opts))
        if not ans:
            continue
        out.append({"type": qtype, "difficulty": diff,
                    "question": mathify(_clean_question(q)),
                    "options": opts, "answer": ans})
    return out


def _parse_duty_fill(ws) -> List[Dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [_s(x) for x in rows[0]]
    diff_i = header.index("难度：") if "难度：" in header else 2
    ans_idx = [i for i, x in enumerate(header) if x.startswith("答案")]
    out: List[Dict] = []
    for row in rows[1:]:
        q = _text(row[0]) if len(row) > 0 else ""
        if not q:
            continue
        answers = [_text(row[i]) for i in ans_idx if i < len(row) and _s(row[i])]
        item: Dict = {"type": "填空",
                      "difficulty": _diff(row[diff_i]) if len(row) > diff_i else "",
                      "question": mathify(_clean_question(q)),
                      "answers": [mathify(a) for a in answers],
                      "orderFixed": False}
        if not answers:
            item["answerMissing"] = True
        out.append(item)
    return out


def _extract_dispimg(text: str) -> Tuple[str, List[str]]:
    ids = re.findall(r'DISPIMG\s*\(\s*"([^"]+)"', text or "", re.I)
    clean = re.sub(r'=?\s*(?:_?xlfn\.?)?\s*DISPIMG\s*\([^)]*\)', "", text or "",
                   flags=re.I).strip()
    return clean, ids


def _parse_duty_qa(ws, anchors: Dict[int, List[str]],
                   cell_map: Dict[str, str]) -> List[Dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    out: List[Dict] = []
    for ridx, row in enumerate(rows[1:], 2):
        raw_q = _text(row[0]) if len(row) > 0 else ""
        raw_a = _text(row[1]) if len(row) > 1 else ""
        if not raw_q:
            continue
        raw_q, q_ids = _extract_dispimg(raw_q)
        raw_a, a_ids = _extract_dispimg(raw_a)
        q = _clean_question(raw_q)
        q_imgs = [cell_map[i] for i in q_ids if i in cell_map]
        a_imgs = [cell_map[i] for i in a_ids if i in cell_map]
        a_imgs += anchors.get(ridx, [])
        if not q:
            # 题干本身只是 WPS DISPIMG 图片、没有可读文字：跳过，避免空题干题
            continue
        qtype = "计算" if _CALC.search(q) else "简答"
        item: Dict = {"type": qtype, "difficulty": "",
                      "question": _maybe_mathify(q, qtype),
                      "answer": _maybe_mathify(raw_a, qtype)}
        if q_imgs:
            item["images"] = list(dict.fromkeys(q_imgs))
        if a_imgs:
            item["answerImages"] = list(dict.fromkeys(a_imgs))
        if not raw_a and not a_imgs:
            item["answerMissing"] = True
        out.append(item)
    return out


def parse_duty(path: str) -> Dict:
    wb = load_workbook(path, data_only=True)
    items: List[Dict] = []
    images: Dict[str, bytes] = {}
    warnings: List[str] = []

    cell_map, cell_blobs = _collect_cell_images(path)
    images.update(cell_blobs)
    for sh in _DUTY_SHEETS:
        if sh not in wb.sheetnames:
            continue
        ws = wb[sh]
        blobs, anchors = _collect_images(ws)
        images.update(blobs)
        if sh == "单选题":
            items += _parse_duty_objective(ws, "单选")
        elif sh == "多选题":
            items += _parse_duty_objective(ws, "多选")
        elif sh == "判断题":
            items += _parse_duty_objective(ws, "判断")
        elif sh == "填空题":
            items += _parse_duty_fill(ws)
        elif sh == "问答题":
            items += _parse_duty_qa(ws, anchors, cell_map)
    wb.close()
    return {"items": items, "images": images, "warnings": warnings}


# ---------------------------------------------------------------------------
# 题库全：客观题解析辅助
# ---------------------------------------------------------------------------

_OPT_START = re.compile(r"^\s*[（(]?\s*([A-L])\s*[）)]?\s*[.．、:：)]\s*(.*)$")
_OPT_INLINE = re.compile(r"(?:(?<=^)|(?<=[\s|]))[（(]?\s*([A-L])\s*[）)]?\s*[.．、:：)]\s*")


def _split_option_text(text: str) -> Tuple[str, "OrderedDict[str, str]"]:
    """
    从一段文本中拆出题干与 A-L 选项。
    支持：
        A. xxx   B、xxx   C) xxx   (D) xxx
        A. xxx | B. xxx | C. xxx D. xxx
        (A)xxx（B)xxx（C)xxx（D)xxx
        A光伏组件 B蓄电池 C逆变器（无标点）
    """
    t = text.replace("\r", "\n").replace("（", "(").replace("）", ")")
    t = t.replace("．", ".").replace("、", ".").replace("：", ":")
    # 把 (A) 统一成 A.
    t = re.sub(r"\(\s*([A-L])\s*\)", r"\n\1. ", t)
    # 把行内 A. / A: / A) 前加换行，避免误伤英文缩写
    t = re.sub(r"(?:^|(?<=[\s\n|]))([A-L])\s*[.:)]\s*", r"\n\1. ", t)
    lines = [ln.strip() for ln in t.split("\n") if ln.strip()]
    opts: "OrderedDict[str, str]" = OrderedDict()
    qparts: List[str] = []
    started = False
    cur: Optional[str] = None
    for ln in lines:
        m = re.match(r"^([A-L])\.\s*(.*)$", ln)
        if m:
            started = True
            cur = m.group(1)
            opts[cur] = m.group(2).strip()
        elif started:
            if cur:
                opts[cur] = (opts[cur] + " " + ln).strip()
        else:
            qparts.append(ln)
    q = " ".join(qparts)
    q = re.sub(r"\s+", " ", q).strip()

    # 兜底：选项没有标点，或 A 选项漏标，如 “A光伏组件 B蓄电池 C逆变器 D控制器”
    need_fallback = (
        len(opts) < 2
        or "A" not in opts
        or any(chr(65 + i) not in opts for i in range(max(ord(k) - 65 for k in opts) + 1))
    )
    if need_fallback:
        t2 = text.replace("\r", "\n").replace("（", "(").replace("）", ")")
        t2 = t2.replace("．", ".").replace("、", ".").replace("：", ":")
        t2 = re.sub(r"\(\s*([A-L])\s*\)", r" \1 ", t2)
        marks = list(re.finditer(
            r"(?:^|(?<=[\s\n|]))([A-L])\s*[.:)]?\s*(?=[^\s\d])", t2))
        good = []
        want = "A"
        for m in marks:
            if m.group(1) == want:
                good.append(m)
                want = chr(ord(want) + 1)
        if len(good) >= 2:
            q2 = t2[:good[0].start()].strip()
            o2: "OrderedDict[str, str]" = OrderedDict()
            for i, m in enumerate(good):
                end = good[i + 1].start() if i + 1 < len(good) else len(t2)
                o2[m.group(1)] = re.sub(r"\s+", " ",
                                        t2[m.end():end]).strip(" .;；,，|")
            q2 = re.sub(r"\s+", " ", q2).strip()
            if len(o2) >= len(opts) or "A" in o2:
                return q2, o2
    return q, opts


def _extract_embedded_answer(q: str) -> Tuple[str, str]:
    """从题干末尾括号中提取答案字母；返回 (清理后的题干, 答案字母)。"""
    # 去掉单元格拼接产生的孤独的 "|"
    q = q.replace("|", " ")
    ms = list(re.finditer(r"[（(]\s*([A-L])\s*[）)]", q))
    if not ms:
        return q, ""
    # 取最后一个只含单个字母的括号作为答案
    m = ms[-1]
    ans = m.group(1)
    q2 = q[:m.start()] + "（ ）" + q[m.end():]
    return q2, ans


def _parse_option_cells(cells: Sequence[str]) -> Tuple[str, "OrderedDict[str, str]"]:
    return _split_option_text("\n".join(c for c in cells if _s(c)))


def _make_choice(q: str, ans: str, opts: "OrderedDict[str, str]",
                 diff: str = "", qtype: str = "单选") -> Optional[Dict]:
    q = _clean_question(q)
    if not q:
        return None
    ans = _answer_letter(ans)
    if not ans:
        return None
    # 只保留有内容的连续选项；若答案字母超出范围则丢弃
    letters = [k for k in "ABCDEFGHIJKL" if k in opts and _s(opts[k])]
    if not letters:
        return None
    max_i = max(ord(k) - 65 for k in letters)
    arr = [mathify(_s(opts.get(chr(65 + i), ""))) for i in range(max_i + 1)]
    if any(not x for x in arr):
        return None
    ans = "".join(c for c in ans if (ord(c) - 65) < len(arr))
    if not ans:
        return None
    if qtype == "单选" and len(ans) != 1:
        qtype = "多选"
    return {"type": qtype, "difficulty": diff,
            "question": mathify(q), "options": arr, "answer": ans}


def _answer_from_lines(lines: Sequence[str], letters: str = "ABCDEFGHIJKL") -> str:
    for ln in lines:
        s = _s(ln).replace("|", "").strip()
        if not s:
            continue
        # 纯字母行 / 字母+括号
        m = re.fullmatch(r"[（(]?\s*([A-L]+)\s*[）)]?\s*[.。]?", s, re.I)
        if m:
            cand = _answer_letter(m.group(1))
            if cand:
                return cand
        # 行首是答案，后面带解释
        m = re.match(r"^\s*[（(]?\s*([A-L]{1,4})\s*[）)]?\s*(?:[.。、,]|\s|$)", s, re.I)
        if m and len(m.group(1)) <= 4:
            cand = _answer_letter(m.group(1))
            if cand:
                return cand
    return ""


# ---------------------------------------------------------------------------
# 题库全：单选题
# ---------------------------------------------------------------------------

def _parse_full_single_structured(rows: List[tuple], lo: int, hi: int) -> List[Dict]:
    """行 2..1285：题干 | 答案 | 选项A..D。"""
    out: List[Dict] = []
    for row in rows[lo - 1:hi]:
        if len(row) < 2:
            continue
        q = _text(row[0])
        ans = _s(row[1])
        if not q or not _answer_letter(ans):
            continue
        cells = [_text(x) for x in row[2:] if _s(x)]
        opts = OrderedDict()
        for i, c in enumerate(cells[:12]):
            opts[chr(65 + i)] = c
        item = _make_choice(q, ans, opts, "", "单选")
        if item:
            out.append(item)
    return out


def _parse_full_single_sixcol(rows: List[tuple], lo: int, hi: int) -> List[Dict]:
    """6 列块：题干 | 选项A | 选项B | 选项C | 选项D | 答案。"""
    out: List[Dict] = []
    for row in rows[lo - 1:hi]:
        if len(row) < 6:
            continue
        q = _text(row[0])
        ans = _s(row[5])
        if not q or not _answer_letter(ans):
            continue
        opts = OrderedDict()
        for i, c in enumerate(row[1:5]):
            if _s(c):
                opts[chr(65 + i)] = _text(c)
        item = _make_choice(q, ans, opts, "", "单选")
        if item:
            out.append(item)
    return out


def _parse_full_numbered(rows: List[tuple], lo: int, hi: int,
                         qtype: str = "单选") -> Tuple[List[Dict], List[str]]:
    """
    编号题块：`1．题干（ B ）。` + 后续选项行。
    选项可能跨行、同行多个、或与题干同一行。
    """
    out: List[Dict] = []
    warnings: List[str] = []
    block: List[str] = []
    qn = 0

    def flush() -> None:
        nonlocal block, qn
        if not block:
            return
        text = "\n".join(block)
        # 先提取题干括号中的答案，再拆选项；否则 (A) 会被误当作选项标记
        q0, ans = _extract_embedded_answer(text)
        q, opts = _split_option_text(q0)
        if not ans:
            # 答案可能在最后一行单独出现
            ans = _answer_from_lines(list(reversed(block)))
        if not q or not ans or len(opts) < 2:
            warnings.append(f"{qtype}编号题解析失败，已跳过：{_s(text)[:60]}")
            block = []
            return
        item = _make_choice(q, ans, opts, "", qtype)
        if item:
            out.append(item)
        else:
            warnings.append(f"{qtype}编号题选项不完整，已跳过：{_s(q)[:60]}")
        block = []

    for i in range(lo - 1, min(hi, len(rows))):
        row = rows[i]
        vals = [_text(x) for x in row if _s(x)]
        if not vals:
            continue
        # 一个单元格里可能塞了多道编号题
        joined = "\n".join(vals)
        # 题号可能出现在同一行中间（如 “... C逆变器 17．...”），统一换行
        joined = re.sub(r"(?<!\d)(\d{1,3})\s*[．、]\s*", r"\n\1．", joined)
        chunks = re.split(r"(?m)(?=^\s*\d{1,3}\s*[．.、]\s*)", joined)
        for chunk in chunks:
            chunk = chunk.strip()
            if not chunk:
                continue
            if re.match(r"^\s*\d{1,3}\s*[．.、]\s*", chunk):
                flush()
                block = [chunk]
                qn += 1
            else:
                if block:
                    block.append(chunk)
    flush()
    return out, warnings


def _parse_full_ocr_choice(rows: List[tuple], lo: int, hi: int,
                           anchors: Dict[int, List[str]],
                           qtype: str = "单选") -> Tuple[List[Dict], List[str]]:
    """
    OCR 块：页标记 `N/120 P` 开始一题，题干/选项可能跨行，`标准答案` 后为答案。
    适用于单选题和判断题。
    """
    out: List[Dict] = []
    warnings: List[str] = []
    blocks: List[Tuple[int, int]] = []
    start = None
    for i in range(lo - 1, min(hi, len(rows))):
        first = _s(rows[i][0]) if rows[i] else ""
        if _PAGE.fullmatch(first):
            if start is not None:
                blocks.append((start, i))
            start = i
    if start is not None:
        blocks.append((start, min(hi, len(rows))))

    for bstart, bend in blocks:
        lines: List[str] = []
        for i in range(bstart, bend):
            row = rows[i]
            for c in row:
                t = _text(c)
                if not t:
                    continue
                if _PAGE.fullmatch(t):
                    continue
                if t in _HEADERS:
                    continue
                if t in ("易", "中", "难", "较易", "较难"):
                    continue
                lines.append(t)
        # 拆分“标准答案”
        before: List[str] = []
        after: List[str] = []
        found = False
        for ln in lines:
            if "标准答案" in ln:
                found = True
                pre, _, post = ln.partition("标准答案")
                if pre.strip():
                    before.append(pre.strip())
                if post.strip():
                    after.append(post.strip())
                continue
            (after if found else before).append(ln)

        if not before:
            continue
        text = "\n".join(before)
        if qtype == "判断":
            q, opts = text, OrderedDict()
        else:
            q, opts = _split_option_text(text)
        ans = _answer_from_lines(after)
        diff = ""
        for ln in lines:
            if ln in ("易", "较易", "中", "中等", "难", "较难"):
                diff = _diff(ln)
                break

        if qtype == "判断":
            jans = _norm_judge(ans)
            if not jans:
                # 答案可能在 after 第一行
                for ln in after:
                    jans = _norm_judge(ln)
                    if jans:
                        break
            if not q or not jans:
                warnings.append(f"判断OCR题解析失败，已跳过：{_s(text)[:60]}")
                continue
            item: Dict = {"type": "判断", "difficulty": diff,
                          "question": mathify(_clean_question(q)),
                          "options": ["√", "×"], "answer": jans}
        else:
            if not q or not ans or len(opts) < 2:
                warnings.append(f"单选OCR题解析失败，已跳过：{_s(text)[:60]}")
                continue
            item = _make_choice(q, ans, opts, diff, qtype)
            if not item:
                warnings.append(f"单选OCR题选项不完整，已跳过：{_s(q)[:60]}")
                continue

        imgs = []
        for r in range(bstart + 1, bend + 1):
            imgs += anchors.get(r, [])
        if imgs:
            item["images"] = list(dict.fromkeys(imgs))
        out.append(item)
    return out, warnings


def _parse_full_single(ws) -> Tuple[List[Dict], Dict[str, bytes], List[str]]:
    rows = list(ws.iter_rows(values_only=True))
    blobs, anchors = _collect_images(ws)
    items: List[Dict] = []
    warnings: List[str] = []

    # 1) 行 2..1285 结构化
    items += _parse_full_single_structured(rows, 2, 1285)
    # 2) 1286..1556 编号题
    got, w = _parse_full_numbered(rows, 1286, 1556, "单选")
    items += got
    warnings += w
    # 3) 1557..2763 OCR
    got, w = _parse_full_ocr_choice(rows, 1557, 2763, anchors, "单选")
    items += got
    warnings += w
    # 4) 2764..2943 6 列块
    items += _parse_full_single_sixcol(rows, 2764, 2943)
    # 5) 2944..3189 编号题
    got, w = _parse_full_numbered(rows, 2944, len(rows), "单选")
    items += got
    warnings += w
    return items, blobs, warnings


# ---------------------------------------------------------------------------
# 题库全：多选题 / 填空题
# ---------------------------------------------------------------------------

def _parse_full_multi(ws) -> Tuple[List[Dict], Dict[str, bytes], List[str]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], {}, []
    header = [_s(x) for x in rows[0]]
    ai = next(i for i, x in enumerate(header) if x.startswith("正确答案"))
    diff_i = header.index("难度：") if "难度：" in header else 2
    opt_i = header.index("选项A") if "选项A" in header else 5
    out: List[Dict] = []
    for row in rows[1:]:
        q = _text(row[0]) if row else ""
        ans = _s(row[ai]) if len(row) > ai else ""
        if not q or not _answer_letter(ans):
            continue
        opts = [mathify(_text(row[i])) for i in range(opt_i, len(row))
                if i < len(row) and _s(row[i])]
        if len(opts) < 2:
            continue
        ans = _answer_letter(ans)
        ans = "".join(c for c in ans if (ord(c) - 65) < len(opts))
        if not ans:
            continue
        out.append({"type": "多选",
                    "difficulty": _diff(row[diff_i]) if len(row) > diff_i else "",
                    "question": mathify(_clean_question(q)),
                    "options": opts, "answer": ans})
    blobs, _ = _collect_images(ws)
    return out, blobs, []


def _parse_full_fill(ws) -> Tuple[List[Dict], Dict[str, bytes], List[str]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], {}, []
    header = [_s(x) for x in rows[0]]
    diff_i = header.index("难度：") if "难度：" in header else 2
    ans_idx = [i for i, x in enumerate(header) if x.startswith("答案")]
    out: List[Dict] = []
    for row in rows[1:]:
        q = _text(row[0]) if row else ""
        if not q:
            continue
        answers = [_text(row[i]) for i in ans_idx if i < len(row) and _s(row[i])]
        item: Dict = {"type": "填空",
                      "difficulty": _diff(row[diff_i]) if len(row) > diff_i else "",
                      "question": mathify(_clean_question(q)),
                      "answers": [mathify(a) for a in answers],
                      "orderFixed": False}
        if not answers:
            item["answerMissing"] = True
        out.append(item)
    blobs, _ = _collect_images(ws)
    return out, blobs, []


# ---------------------------------------------------------------------------
# 题库全：判断题
# ---------------------------------------------------------------------------

def _parse_full_judge(ws) -> Tuple[List[Dict], Dict[str, bytes], List[str]]:
    rows = list(ws.iter_rows(values_only=True))
    blobs, anchors = _collect_images(ws)
    items: List[Dict] = []
    warnings: List[str] = []

    # 1) 结构化 2..1491
    for row in rows[1:1491]:
        q = _text(row[0]) if row else ""
        ans = _norm_judge(row[4]) if len(row) > 4 else ""
        if not q or not ans:
            continue
        items.append({"type": "判断",
                      "difficulty": _diff(row[2]) if len(row) > 2 else "",
                      "question": mathify(_clean_question(q)),
                      "options": ["√", "×"], "answer": ans})

    # 2) 1492..1617 编号题（答案在末尾括号，可能跨行）
    text = "\n".join(
        _text(x) for row in rows[1491:1617] for x in row if _s(x)
    )
    chunks = re.split(r"(?m)(?=^\s*\d{1,3}\s*[．.、]\s*)", text)
    for chunk in chunks:
        chunk = chunk.replace("|", "").strip()
        if not chunk:
            continue
        m = re.match(r"^\s*\d{1,3}\s*[．.、]\s*(.*)$", chunk, re.S)
        if not m:
            continue
        body = m.group(1)
        # 末尾括号答案
        am = list(re.finditer(r"[（(]\s*([√×xX对错])\s*[）)]", body))
        if not am:
            warnings.append(f"判断编号题缺答案，已跳过：{_s(body)[:50]}")
            continue
        ans = _norm_judge(am[-1].group(1))
        q = body[:am[-1].start()] + body[am[-1].end():]
        if not q.strip() or not ans:
            continue
        items.append({"type": "判断", "difficulty": "",
                      "question": mathify(_clean_question(q)),
                      "options": ["√", "×"], "answer": ans})

    # 3) 1618..2401 OCR
    got, w = _parse_full_ocr_choice(rows, 1618, 2401, anchors, "判断")
    items += got
    warnings += w

    # 4) 2402..2581 两列块：题干 | 对/错
    for row in rows[2401:2581]:
        q = _text(row[0]) if row else ""
        ans = _norm_judge(row[1]) if len(row) > 1 else ""
        if not q or not ans:
            continue
        items.append({"type": "判断", "difficulty": "",
                      "question": mathify(_clean_question(q)),
                      "options": ["√", "×"], "answer": ans})

    # 5) 2582..末尾 编号题（答案在末尾括号，可能跨行）
    text = "\n".join(
        _text(x) for row in rows[2581:] for x in row if _s(x)
    )
    chunks = re.split(r"(?m)(?=^\s*\d{1,3}\s*[．.、]\s*)", text)
    for chunk in chunks:
        chunk = chunk.replace("|", "").strip()
        if not chunk:
            continue
        m = re.match(r"^\s*\d{1,3}\s*[．.、]\s*(.*)$", chunk, re.S)
        if not m:
            continue
        body = m.group(1)
        am = list(re.finditer(r"[（(]\s*([√×xX对错])\s*[）)]", body))
        if not am:
            warnings.append(f"判断编号题缺答案，已跳过：{_s(body)[:50]}")
            continue
        ans = _norm_judge(am[-1].group(1))
        q = body[:am[-1].start()] + body[am[-1].end():]
        if not q.strip() or not ans:
            continue
        items.append({"type": "判断", "difficulty": "",
                      "question": mathify(_clean_question(q)),
                      "options": ["√", "×"], "answer": ans})

    return items, blobs, warnings


# ---------------------------------------------------------------------------
# 题库全：简答 / 论述（三块结构）
# ---------------------------------------------------------------------------

def _numbered_qa(lines: List[str]) -> List[Dict]:
    """把“1．题干 / 答：...”文本块解析为简答题。"""
    out: List[Dict] = []
    cur_q: Optional[str] = None
    cur_a: List[str] = []

    def flush() -> None:
        nonlocal cur_q, cur_a
        if not cur_q:
            cur_a = []
            return
        answer = "\n".join(x for x in cur_a if x.strip())
        q, answer = _fix_subjective_qa(_clean_question(cur_q), answer)
        out.append({"type": "简答", "difficulty": "",
                    "question": mathify(q),
                    "answer": mathify(answer)})
        cur_q, cur_a = None, []

    for ln in lines:
        m = re.match(r"^\s*\d{1,3}\s*[．.、]\s*(.*)$", ln)
        if m:
            flush()
            cur_q = m.group(1)
        elif cur_q is not None:
            cur_a.append(ln)
    flush()
    return out


def _ocr_subjective(rows: List[tuple], lo: int, hi: int,
                    anchors: Dict[int, List[str]]) -> List[Dict]:
    """页标记开始的简答/论述/计算/绘图块。"""
    out: List[Dict] = []
    blocks: List[Tuple[int, int]] = []
    start = None
    for i in range(lo - 1, min(hi, len(rows))):
        first = _s(rows[i][0]) if rows[i] else ""
        if _PAGE.fullmatch(first):
            if start is not None:
                blocks.append((start, i))
            start = i
    if start is not None:
        blocks.append((start, min(hi, len(rows))))

    for bstart, bend in blocks:
        raw: List[str] = []
        for i in range(bstart, bend):
            for c in rows[i]:
                t = _text(c)
                if t and not _PAGE.fullmatch(t):
                    raw.append(t)
        if not raw:
            continue

        # 题型：简答题 / 论述题 / 计算题 / 绘图题，可能和难度连写
        qtype = "简答"
        diff = ""
        cleaned: List[str] = []
        for ln in raw:
            s = ln.strip()
            if s in _HEADERS:
                if s.endswith("题"):
                    qtype = s[:-1] if s != "问答题" else "简答"
                continue
            m = re.match(r"^(简答题|论述题|计算题|绘图题)\s*(易|较易|中|中等|难|较难)?$", s)
            if m:
                qtype = m.group(1)[:-1]
                if m.group(2):
                    diff = _diff(m.group(2))
                continue
            if s in ("易", "较易", "中", "中等", "难", "较难"):
                diff = _diff(s)
                continue
            cleaned.append(s)

        before: List[str] = []
        after: List[str] = []
        found = False
        for ln in cleaned:
            if "标准答案" in ln:
                found = True
                pre, _, post = ln.partition("标准答案")
                if pre.strip():
                    before.append(pre.strip())
                if post.strip():
                    after.append(post.strip())
                continue
            (after if found else before).append(ln)

        q = "\n".join(before).strip()
        a = "\n".join(after).strip()
        if not q:
            continue
        q = _clean_question(q)
        imgs = []
        for r in range(bstart + 1, bend + 1):
            imgs += anchors.get(r, [])
        imgs = list(dict.fromkeys(imgs))

        if qtype == "绘图":
            # 绘图题：答案图片必保留；若原表已有文字说明也一并保留
            if not imgs:
                continue
            item: Dict = {"type": "绘图", "difficulty": diff,
                          "question": q, "answer": a,
                          "answerImages": imgs}
            out.append(item)
            continue

        item: Dict = {"type": qtype, "difficulty": diff,
                      "question": _maybe_mathify(q, qtype),
                      "answer": _maybe_mathify(a, qtype)}
        if imgs:
            if qtype == "计算":
                # 计算题：有图则直接用原图（图题放题干图，解法图放答案图）
                if re.search(r"如图|图所示|图示|图中|下图|上图", q):
                    item["images"] = imgs
                else:
                    item["answerImages"] = imgs
            else:
                if re.search(r"如图|图所示|图示|图中|下图|上图", q):
                    item["images"] = imgs
                else:
                    item["answerImages"] = imgs
        if not a and not imgs:
            item["answerMissing"] = True
        out.append(item)
    return out


def _parse_full_subjective(ws, sheet: str) -> Tuple[List[Dict], Dict[str, bytes], List[str]]:
    rows = list(ws.iter_rows(values_only=True))
    blobs, anchors = _collect_images(ws)
    out: List[Dict] = []
    warnings: List[str] = []

    if sheet == "简答":
        # 1) 2..619 编号简答
        lines = [_text(row[0]) for row in rows[1:619] if row and _s(row[0])]
        out += _numbered_qa(lines)
        # 2) 620..1273 OCR 简答/论述
        out += _ocr_subjective(rows, 620, 1273, anchors)
        # 3) 1274..1443 两列 Q/A
        for row in rows[1273:1443]:
            q = _text(row[0]) if row else ""
            a = _text(row[1]) if len(row) > 1 else ""
            if not q or q in _HEADERS:
                continue
            if not a:
                continue
            out.append({"type": "简答", "difficulty": "",
                        "question": mathify(_clean_question(q)),
                        "answer": mathify(a)})
        # 4) 1444..末尾 编号简答（答案可能内联）
        lines = [_text(row[0]) for row in rows[1443:] if row and _s(row[0])]
        out += _numbered_qa(lines)
    else:
        # 计算 / 绘图：仅保留带图片的块
        got = _ocr_subjective(rows, 2, len(rows), anchors)
        for it in got:
            if it.get("images") or it.get("answerImages"):
                out.append(it)
    return out, blobs, warnings


# ---------------------------------------------------------------------------
# 融合与去重
# ---------------------------------------------------------------------------

def _item_key(it: Dict) -> tuple:
    t = it.get("type", "")
    q = _norm(it.get("question", ""))
    if t in ("单选", "多选"):
        opts = tuple(_norm(x) for x in it.get("options", []))
        return (t, q, opts, _norm(it.get("answer", "")))
    if t == "判断":
        return (t, q, _norm(it.get("answer", "")))
    if t == "填空":
        return (t, q, tuple(_norm(x) for x in it.get("answers", [])))
    if t == "计算":
        # 计算题按题干+文字答案去重；若同题在别处带图，合并时保留图片
        return (t, q, _norm(it.get("answer", "")))
    if t == "绘图":
        # 绘图题以图片作为答案；图片不同视为不同题，避免把多张答案图合并
        imgs = tuple(sorted((it.get("images") or []) + (it.get("answerImages") or [])))
        return (t, q, _norm(it.get("answer", "")), imgs)
    return (t, q, _norm(it.get("answer", "")))


def _merge_item(old: Dict, new: Dict) -> None:
    # 答案更完整者优先
    if len(_s(new.get("answer"))) > len(_s(old.get("answer"))):
        old["answer"] = new.get("answer", old.get("answer", ""))
    if not old.get("difficulty") and new.get("difficulty"):
        old["difficulty"] = new["difficulty"]
    for k in ("images", "answerImages"):
        if new.get(k):
            merged = list(dict.fromkeys(list(old.get(k, [])) + list(new[k])))
            old[k] = merged
    if old.get("answerMissing") and not new.get("answerMissing"):
        old.pop("answerMissing", None)
    if new.get("answer") and not old.get("answer"):
        old.pop("answerMissing", None)


def dedupe(items: Iterable[Dict]) -> List[Dict]:
    seen: "OrderedDict[tuple, Dict]" = OrderedDict()
    for it in items:
        k = _item_key(it)
        if k in seen:
            _merge_item(seen[k], it)
        else:
            seen[k] = dict(it)
    return list(seen.values())


def parse_merged(duty_path: str, full_path: str) -> Dict:
    duty = parse_duty(duty_path)
    full_wb = load_workbook(full_path, data_only=True)
    full_items: List[Dict] = []
    images: Dict[str, bytes] = dict(duty["images"])
    warnings: List[str] = list(duty["warnings"])

    if "单选题" in full_wb.sheetnames:
        got, blobs, w = _parse_full_single(full_wb["单选题"])
        full_items += got
        images.update(blobs)
        warnings += w
    if "多选题" in full_wb.sheetnames:
        got, blobs, w = _parse_full_multi(full_wb["多选题"])
        full_items += got
        images.update(blobs)
        warnings += w
    if "判断题" in full_wb.sheetnames:
        got, blobs, w = _parse_full_judge(full_wb["判断题"])
        full_items += got
        images.update(blobs)
        warnings += w
    if "填空题" in full_wb.sheetnames:
        got, blobs, w = _parse_full_fill(full_wb["填空题"])
        full_items += got
        images.update(blobs)
        warnings += w
    for sh in ("简答", "计算", "绘图"):
        if sh in full_wb.sheetnames:
            got, blobs, w = _parse_full_subjective(full_wb[sh], sh)
            full_items += got
            images.update(blobs)
            warnings += w
    full_wb.close()

    items = dedupe(list(duty["items"]) + full_items)
    for it in items:
        t = it.get("type")
        if t in ("简答", "论述", "名词解释"):
            q, a = _fix_subjective_qa(it.get("question", ""), it.get("answer", ""))
            it["question"] = mathify(q)
            if it.get("answer") or a:
                it["answer"] = mathify(a)
        # 计算/绘图题保留原始文字，不做 LaTeX 转换
    for i, it in enumerate(items, 1):
        it["id"] = i
    return {"items": items, "images": images, "warnings": warnings}